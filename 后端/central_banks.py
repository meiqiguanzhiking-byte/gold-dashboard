"""已报告官方黄金储备月度净变化；显式范围、原表空白与CSV逐月核验。"""
import hashlib
import json
import math
import re
from zipfile import BadZipFile
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from config import CENTRAL_BANK_CHANGES_FILE

SOURCE = "全球央行购金.csv"
SCOPE_FILE = Path(__file__).with_name("central_bank_scope.json")
NOTICE = "基于已报告数据，可能随补报修订；不代表完整全球总量。"


def read_reported_months(path):
    """只读原始Monthly表；保留数值0，原表空白保持缺失，不用CSV的补零判断报告。"""
    with Path(path).open("rb") as stream:
        book = load_workbook(stream, read_only=True, data_only=True)
        try:
            rows = iter(book["Monthly"].values)
            header = next(row for row in rows if "Country" in row and "Comments" in row)
            col = header.index("Country")
            dates = [(i, d.strftime("%Y-%m")) for i, d in enumerate(header) if isinstance(d, (date, datetime))]
            result, entities = {d: {} for _, d in dates}, set()
            if len(result) != len(dates) or not dates:
                raise ValueError("原表月份重复或缺失")
            for row in rows:
                name = str(row[col] or "").strip()
                if not name:
                    continue
                if name in entities:
                    raise ValueError("原表机构重复")
                entities.add(name)
                for i, month in dates:
                    v = row[i]
                    if v is None or isinstance(v, str) and not v.strip():
                        result[month][name] = None
                    elif isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v):
                        result[month][name] = v
                    else:
                        raise ValueError("原表包含非有限数字或未识别文本")
            return result, entities
        finally:
            book.close()


def add_central_bank_changes(builder, evidence_file=CENTRAL_BANK_CHANGES_FILE):
    scope = json.loads(SCOPE_FILE.read_text(encoding="utf-8"))
    included, excluded = scope["included_entities"], scope["excluded_entities"]
    if len(set(included)) != len(included) or set(included) & set(excluded):
        raise ValueError("机构纳入与排除名单冲突")
    table = builder.tables.get((SOURCE, None))
    frame = table.frame.copy() if table else pd.DataFrame(columns=["时间"])
    if "时间" not in frame:
        frame = pd.DataFrame(columns=["时间", *frame.columns])
    columns = [c for c in frame if c != "时间"]
    schema_ok = set(columns) == set(included) | set(excluded)
    notes = [NOTICE, "原值为储备月度净变化，不再差分；正值合计为买入量，负值合计为卖出量，二者相加为净买入。",
             "买入/卖出是机构净变化正负分组，不是交易流水；原表数值0计为有值，空白不计为报告。"]
    try:
        reported, entities = read_reported_months(evidence_file)
        schema_ok = schema_ok and set(columns) == entities
    except (OSError, ValueError, KeyError, StopIteration, BadZipFile, TypeError, IndexError):
        reported = {}
        notes.append("数据暂缺：WGC原始Monthly表不可用或结构异常，无法核实CSV历史补零和报告覆盖。")
    if not schema_ok:
        notes.append("机构字段发生变化，待核对纳入/排除名单；相关汇总不计算。")
    sums = {"purchases": [], "sales": [], "net": [], "count": []}
    # 未通过逐月原表核验时，国家序列也不把CSV补零误当有效报告。
    restored = frame.copy()
    for c in columns:
        restored[c] = pd.Series(float("nan"), index=frame.index)
    bad_months = []
    for idx, row in frame.iterrows():
        month = str(row["时间"])[:7]
        original = reported.get(month)
        verified = schema_ok and original is not None
        if verified:
            for c in columns:
                try:
                    actual = float(row[c])
                except (ValueError, TypeError):
                    actual = float("nan")
                expected = original[c]
                if expected is None:
                    matches = pd.isna(actual) or actual == 0
                else:
                    matches = math.isfinite(actual) and math.isclose(actual, expected, rel_tol=1e-10, abs_tol=1e-9)
                if not matches:
                    verified = False
                    break
        if not verified:
            bad_months.append(month)
            for values in sums.values():
                values.append(None)
            continue
        for c in columns:
            restored.at[idx, c] = original[c]
        values = [original[c] for c in included if original[c] is not None]
        positive = math.fsum(v for v in values if v > 0)
        negative = math.fsum(v for v in values if v < 0)
        sums["purchases"].append(positive if values else None)
        sums["sales"].append(negative if values else None)
        sums["net"].append(positive + negative if values else None)
        sums["count"].append(len(values))
    if bad_months:
        notes.append(f"{len(bad_months)}个月份缺少匹配的原始报告或与CSV不一致，相关数据保持缺失。")
    for country in columns:
        slug = re.sub(r"[^a-z0-9]+", "_", country.lower()).strip("_")
        suffix = hashlib.sha1(country.encode("utf-8")).hexdigest()[:8]
        builder.add("central-bank", f"central_bank_country_{slug}_{suffix}", f"{country} 黄金储备月度净变化",
                    "吨", "月度", SOURCE, column=country, frame=restored,
                    warnings=notes, country=country, aggregation_included=country in included,
                    entity_type="纳入官方储备实体" if country in included else "排除：" + excluded.get(country, "待核实"),
                    source_reference=scope["source_reference"])
    for field, key, name, unit in [
        ("purchases", "global_central_bank_gross_purchases", "月买入量（已报告正变化合计）", "吨"),
        ("sales", "global_central_bank_gross_sales", "月卖出量（已报告负变化合计）", "吨"),
        ("net", "global_central_bank_net_purchases", "全球黄金储备净变化合计（已报告）", "吨"),
        ("count", "global_central_bank_reporting_count", "月度有效报告机构数", "家"),
    ]:
        builder.add("central-bank", key, name, unit, "月度", SOURCE, frame=frame, values=sums[field], warnings=notes,
                    source_reference=scope["source_reference"], scope_entity_count=len(included),
                    aggregation_scope=scope if field == "net" else None,
                    coverage_basis="已报告数据，非完整全球总量", reporting_evidence="WGC Changes / Monthly",
                    method="Σ正变化；Σ负变化；净变化＝正变化合计＋负变化合计；报告数包含原表真实0，排除原表空白")
